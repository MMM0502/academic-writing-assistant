# -*- coding: utf-8 -*-
"""从学校论文模板（.docx）中抽取可执行的排版规则。

设计取舍：
- 只做 DOCX。学校模板绝大多数是 Word，PDF 抽取误差太大，先不承诺。
- 取值优先级：显式样式 → docDefaults（文档默认字体）→ 段落特征众数。
  真实模板里"正文样式没字体、字体只写在 docDefaults"非常常见，必须覆盖。
- 行距同时输出倍数（给现有渲染器）与固定磅值（Word 里常见的"固定值 20 磅"）。
- 不猜内容结构，只抽"排版规则"；章节角色识别留给应用器。
"""
from __future__ import annotations

from collections import Counter
from io import BytesIO
from typing import Any

try:  # python-docx 是可选依赖
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Length

    HAS_DOCX = True
except ImportError:  # pragma: no cover
    HAS_DOCX = False


FONT_MAP = {
    "宋体": "SimSun",
    "simsun": "SimSun",
    "黑体": "SimHei",
    "simhei": "SimHei",
    "楷体": "KaiTi",
    "楷体_gb2312": "KaiTi",
    "kaiti": "KaiTi",
    "仿宋": "FangSong",
    "仿体": "FangSong",
    "fangsong": "FangSong",
    "微软雅黑": "Microsoft YaHei",
    "times new roman": "Times New Roman",
}

BODY_STYLE_NAMES = ["normal", "正文", "body text", "body"]
HEADING_STYLE_NAMES = {
    1: ["heading 1", "标题 1", "标题1"],
    2: ["heading 2", "标题 2", "标题2"],
    3: ["heading 3", "标题 3", "标题3"],
}

SPACING_MULTIPLES = {240: "1.0", 288: "1.2", 360: "1.5", 480: "2.0"}


def _map_font(name: str | None) -> str | None:
    if not name:
        return None
    key = name.strip().lower()
    return FONT_MAP.get(key) or FONT_MAP.get(name.strip()) or name.strip()


def _rfonts(element) -> tuple[str | None, str | None]:
    """从 rPr 里读 ascii / eastAsia 字体。"""
    try:
        rpr = element.find(qn("w:rPr"))
        if rpr is None:
            return None, None
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            return None, None
        return rfonts.get(qn("w:ascii")), rfonts.get(qn("w:eastAsia"))
    except Exception:
        return None, None


def _style_fonts(style) -> tuple[str | None, str | None]:
    ascii_name, east = _rfonts(style.element)
    ascii_name = ascii_name or (style.font.name if style.font is not None else None)
    return ascii_name, east


def _doc_defaults_fonts(doc) -> tuple[str | None, str | None]:
    """docDefaults/rPrDefault 里的字体：很多模板只在这里定义正文宋体。"""
    try:
        styles_el = doc.styles.element
        defaults = styles_el.find(qn("w:docDefaults"))
        if defaults is None:
            return None, None
        rpr_default = defaults.find(qn("w:rPrDefault"))
        if rpr_default is None:
            return None, None
        rpr = rpr_default.find(qn("w:rPr"))
        if rpr is None:
            return None, None
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            return None, None
        return rfonts.get(qn("w:ascii")), rfonts.get(qn("w:eastAsia"))
    except Exception:
        return None, None


def _style_size_pt(style) -> float | None:
    try:
        if style.font is not None and style.font.size is not None:
            return round(style.font.size.pt, 1)
    except Exception:
        pass
    return None


def _alignment_name(value) -> str | None:
    """把 python-docx 的枚举（如 'JUSTIFY (3)'）压成规则里的规范值。"""
    if value is None:
        return None
    text = str(value).upper()
    if "JUSTIFY" in text or "BOTH" in text:
        return "both"
    if "CENTER" in text:
        return "center"
    if "RIGHT" in text:
        return "right"
    if "LEFT" in text:
        return "left"
    return None


def _style_alignment(style) -> str | None:
    try:
        fmt = style.paragraph_format
        if fmt is None:
            return None
        return _alignment_name(fmt.alignment)
    except Exception:
        return None


def _line_spacing(style) -> dict[str, str]:
    """返回 {'line_spacing': '1.5'} 或 {'line_spacing_pt': '20'}。"""
    try:
        fmt = style.paragraph_format
        if fmt is None or fmt.line_spacing is None:
            return {}
        value = fmt.line_spacing
        if isinstance(value, Length):
            return {"line_spacing_pt": f"{value.pt:g}"}
        multiple = float(value)
        return {"line_spacing": SPACING_MULTIPLES.get(round(multiple * 240), f"{multiple:g}")}
    except Exception:
        return {}


def _paragraph_line_spacing_mode(doc) -> dict[str, str]:
    """样式没写行距时，按正文段落的 w:spacing 取众数。"""
    counter: Counter = Counter()
    for p in doc.paragraphs:
        if len(_paragraph_text(p)) < 40:
            continue
        try:
            ppr = p._p.pPr
            if ppr is None:
                continue
            spacing = ppr.find(qn("w:spacing"))
            if spacing is None:
                continue
            line = spacing.get(qn("w:line"))
            rule = spacing.get(qn("w:lineRule")) or "auto"
            if line:
                counter[(int(line), rule)] += 1
        except Exception:
            continue
    if not counter:
        return {}
    (line, rule), _count = counter.most_common(1)[0]
    if rule in ("exact", "atLeast"):
        return {"line_spacing_pt": f"{line / 20:g}"}
    return {"line_spacing": SPACING_MULTIPLES.get(line, f"{line / 240:g}")}


def _style_first_indent_chars(style, size_pt: float | None) -> int | None:
    try:
        fmt = style.paragraph_format
        if fmt is None or fmt.first_line_indent is None:
            return None
        pt = fmt.first_line_indent.pt
        if not size_pt:
            return 2 if pt > 0 else 0
        return max(0, round(pt / size_pt))
    except Exception:
        return None


def _find_style(doc, names: list[str]):
    wanted = {name.lower() for name in names}
    for style in doc.styles:
        try:
            if style.name and style.name.strip().lower() in wanted:
                return style
        except Exception:
            continue
    return None


def _paragraph_text(p) -> str:
    try:
        return (p.text or "").strip()
    except Exception:
        return ""


def _paragraph_run_stats(doc) -> Counter:
    """统计正文段落的（字体, 字号）众数，作为样式缺失时的兜底。"""
    counter: Counter = Counter()
    for p in doc.paragraphs:
        text = _paragraph_text(p)
        if len(text) < 40:
            continue
        for run in p.runs:
            if not run.text.strip():
                continue
            east = None
            try:
                rpr = run._element.rPr
                if rpr is not None and rpr.rFonts is not None:
                    east = rpr.rFonts.get(qn("w:eastAsia"))
            except Exception:
                east = None
            size = round(run.font.size.pt, 1) if run.font is not None and run.font.size is not None else None
            if east or size:
                counter[(east, size)] += len(run.text)
    return counter


def _paragraph_sizes(doc, min_len: int = 0, max_len: int = 99999) -> Counter:
    counter: Counter = Counter()
    for p in doc.paragraphs:
        text = _paragraph_text(p)
        if not (min_len <= len(text) <= max_len):
            continue
        for run in p.runs:
            if run.font is not None and run.font.size is not None:
                counter[round(run.font.size.pt, 1)] += 1
    return counter


def _heading_candidates(doc, body_size: float | None) -> dict[int, list]:
    """按编号形态收集疑似标题段落，只保留明显大于正文的字号。"""
    buckets: dict[int, list] = {1: [], 2: [], 3: []}
    for p in doc.paragraphs:
        text = _paragraph_text(p)
        if not text or len(text) > 60:
            continue
        level = None
        if text.startswith("第") and ("章" in text[:8] or "节" in text[:8]):
            level = 1
        elif text[:2] in {f"{c}、" for c in "一二三四五六七八九十"}:
            level = 1
        elif text[0].isdigit() and "." in text[:6]:
            level = 3 if text.count(".") >= 2 else 2
        elif text.startswith("（") and "）" in text[:5]:
            level = 3
        if not level:
            continue
        for run in p.runs:
            size = round(run.font.size.pt, 1) if run.font is not None and run.font.size is not None else None
            if size and (body_size is None or size > body_size):
                buckets[level].append(size)
                break
    return buckets


def _footer_page_number(doc) -> tuple[str, str]:
    """返回 (是否检测到页码字段, 对齐方式)。"""
    try:
        for section in doc.sections:
            for part in (section.footer, section.first_page_footer, section.even_page_footer):
                if part is None:
                    continue
                xml = part._element.xml
                if "PAGE" not in xml:
                    continue
                for p in part.paragraphs:
                    if p.paragraph_format is not None and p.paragraph_format.alignment is not None:
                        return "yes", _alignment_name(p.paragraph_format.alignment) or "center"
                return "yes", "center"
    except Exception:
        pass
    return "no", "center"


def extract_docx_rules(data: bytes) -> dict[str, Any]:
    """抽取模板规则。返回 {rules, meta}。"""
    if not HAS_DOCX:
        raise ValueError("缺少 python-docx，无法解析 Word 模板。请先 pip install python-docx。")

    document = Document(BytesIO(data))
    warnings: list[str] = []
    derived: dict[str, str] = {}
    rules: dict[str, Any] = {}

    # ---------- 页面设置 ----------
    try:
        section = document.sections[0]
        rules["margin_top"] = f"{section.top_margin.cm:.2f}"
        rules["margin_bottom"] = f"{section.bottom_margin.cm:.2f}"
        rules["margin_left"] = f"{section.left_margin.cm:.2f}"
        rules["margin_right"] = f"{section.right_margin.cm:.2f}"
        rules["page_width_cm"] = f"{section.page_width.cm:.2f}"
        rules["page_height_cm"] = f"{section.page_height.cm:.2f}"
        derived["margins"] = "section"
    except Exception as exc:
        warnings.append(f"页面设置读取失败：{exc}")

    # ---------- 正文 ----------
    body_style = _find_style(document, BODY_STYLE_NAMES)
    size_pt = _style_size_pt(body_style) if body_style is not None else None
    style_ascii, style_east = _style_fonts(body_style) if body_style is not None else (None, None)
    default_ascii, default_east = _doc_defaults_fonts(document)
    east = style_east or style_ascii or default_east or default_ascii

    source = "style" if (size_pt and east) else "partial"
    if not east:
        stats = _paragraph_run_stats(document)
        if stats:
            (fallback_east, _fallback_size), _count = stats.most_common(1)[0]
            east = fallback_east
    if not size_pt:
        sizes = _paragraph_sizes(document, min_len=40)
        if sizes:
            size_pt = sizes.most_common(1)[0][0]
    if east:
        rules["font_family"] = _map_font(east)
    if size_pt:
        rules["font_size"] = f"{size_pt:g}pt"
    derived["body"] = source
    if source != "style":
        warnings.append("正文样式未完整定义，已结合文档默认字体与正文段落众数推断。")
    if not east:
        warnings.append("未识别出正文字体。")
    if not size_pt:
        warnings.append("未识别出正文字号。")

    spacing = _line_spacing(body_style) if body_style is not None else {}
    if not spacing:
        spacing = _paragraph_line_spacing_mode(document)
        if spacing:
            warnings.append("正文样式未定义行距，已按正文段落众数推断。")
    if not spacing:
        spacing = {"line_spacing": "1.5"}
        warnings.append("未识别出行距，按 1.5 倍兜底。")
    rules.update(spacing)

    indent_chars = _style_first_indent_chars(body_style, size_pt) if body_style is not None else None
    if indent_chars is None:
        rules["first_line_indent"] = "2chars"
        warnings.append("未识别出首行缩进，按 2 字符兜底。")
    else:
        rules["first_line_indent"] = "2chars" if indent_chars >= 1 else "none"

    alignment = _style_alignment(body_style) if body_style is not None else None
    rules["body_alignment"] = alignment or "both"

    # ---------- 标题 ----------
    candidates = _heading_candidates(document, size_pt)
    for level in (1, 2, 3):
        style = _find_style(document, HEADING_STYLE_NAMES[level])
        h_east = None
        h_size = None
        heading_source = None
        if style is not None:
            h_ascii, h_east_style = _style_fonts(style)
            h_east = h_east_style or h_ascii
            h_size = _style_size_pt(style)
            heading_source = "style" if (h_east and h_size) else "partial-style"
        if (not h_size) and candidates[level]:
            h_size = Counter(candidates[level]).most_common(1)[0][0]
            heading_source = heading_source or "paragraph-mode"
        if not h_east:
            h_east = default_east or default_ascii or "SimHei"
            if heading_source == "style":
                heading_source = "partial-style"
        rules[f"heading{level}_font"] = _map_font(h_east) or "SimHei"
        if h_size:
            rules[f"heading{level}_size"] = f"{h_size:g}pt"
        else:
            default = {1: 16, 2: 14, 3: 12}[level]
            rules[f"heading{level}_size"] = f"{default}pt"
            warnings.append(f"标题 {level} 字号未识别，按 {default}pt 兜底。")
        derived[f"heading{level}"] = heading_source or "fallback"

    # ---------- 参考文献 ----------
    ref_style = None
    for style in document.styles:
        try:
            if style.name and "参考" in style.name:
                ref_style = style
                break
        except Exception:
            continue
    ref_paragraph = None
    for index, p in enumerate(document.paragraphs):
        if _paragraph_text(p) == "参考文献":
            ref_paragraph = document.paragraphs[index + 1] if index + 1 < len(document.paragraphs) else None
            break
    ref_source = ref_style or (ref_paragraph.style if ref_paragraph is not None else None)
    if ref_source is not None:
        ref_ascii, ref_east = _style_fonts(ref_source)
        rules["ref_font"] = _map_font(ref_east or ref_ascii or east or "SimSun")
        ref_size = _style_size_pt(ref_source)
        if ref_size:
            rules["ref_size"] = f"{ref_size:g}pt"
    else:
        rules["ref_font"] = _map_font(east) or "SimSun"
        warnings.append("未识别出参考文献样式，按正文字体兜底。")

    # ---------- 页眉页码 ----------
    has_page, page_align = _footer_page_number(document)
    rules["page_number"] = has_page
    rules["page_number_pos"] = page_align
    if has_page == "no":
        warnings.append("未检测到页码字段。")

    # ---------- 置信度 ----------
    explicit = sum(1 for value in derived.values() if value == "style")
    confidence = 0.5 + 0.12 * explicit
    if derived.get("body") != "style":
        confidence -= 0.08
    if len(warnings) >= 4:
        confidence -= 0.08
    confidence = max(0.3, min(0.95, round(confidence, 2)))

    fields = {
        "margins": derived.get("margins") == "section",
        "body": derived.get("body") == "style",
        "headings": all(derived.get(f"heading{level}") == "style" for level in (1, 2, 3)),
    }

    first_text = ""
    for p in document.paragraphs:
        if _paragraph_text(p):
            first_text = _paragraph_text(p)
            break

    return {
        "rules": rules,
        "meta": {
            "source": "docx",
            "confidence": confidence,
            "fields": fields,
            "warnings": warnings,
            "paragraphs": len(document.paragraphs),
            "title_guess": first_text,
            "derived": derived,
        },
    }
